import html
import io
import json
import re
from typing import List, Tuple
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from app.core.exceptions import ResAIException


def parse_markdown_table(lines: List[str]) -> Tuple[List[List[str]], int]:
    """Parses contiguous markdown table lines into a 2D matrix of cells and returns rows & lines consumed."""
    rows = []
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if not line.startswith("|") and not ("|" in line and not line.startswith("#")):
            break
        # Skip separator line |---|---|
        if re.match(r"^\|?\s*[-:]+[-| :]*$", line):
            idx += 1
            continue
        # Split cells
        cells = [c.strip() for c in line.split("|")]
        # Trim leading/trailing empty cells from outer pipes
        if cells and cells[0] == "":
            cells.pop(0)
        if cells and cells[-1] == "":
            cells.pop()
        if cells:
            rows.append(cells)
        idx += 1
    return rows, idx


class ExportService:
    @staticmethod
    def export_markdown(markdown_content: str) -> bytes:
        return markdown_content.encode("utf-8")

    @staticmethod
    def export_json(data: dict) -> bytes:
        return json.dumps(data, indent=2, default=str).encode("utf-8")

    @staticmethod
    def export_csv_sources_and_evidence(sources: list, evidence: list) -> bytes:
        df_evidence = pd.DataFrame(evidence)
        if df_evidence.empty:
            df_evidence = pd.DataFrame([{"info": "No evidence items recorded"}])
        
        output = io.StringIO()
        df_evidence.to_csv(output, index=False)
        return output.getvalue().encode("utf-8")

    @staticmethod
    def export_html(title: str, markdown_content: str) -> bytes:
        clean_doc_title = title.replace("Research Report: ", "").strip()
        
        # Convert markdown lines to HTML
        lines = markdown_content.split("\n")
        body_parts = []
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue

            # Check if start of table
            if line.startswith("|") or ("|" in line and i + 1 < len(lines) and re.match(r"^\|?\s*[-:]+[-| :]*$", lines[i+1].strip())):
                table_rows, consumed = parse_markdown_table(lines[i:])
                if table_rows:
                    table_html = ['<div class="table-container"><table>']
                    for r_idx, row in enumerate(table_rows):
                        tag = "th" if r_idx == 0 else "td"
                        row_html = "<tr>"
                        for cell in row:
                            formatted_cell = html.escape(cell)
                            formatted_cell = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", formatted_cell)
                            formatted_cell = re.sub(
                                r"\[(.*?)\]\((https?://[^\s\)]+)\)",
                                r'<a href="\2" target="_blank" rel="noopener noreferrer">\1</a>',
                                formatted_cell,
                            )
                            row_html += f"<{tag}>{formatted_cell}</{tag}>"
                        row_html += "</tr>"
                        table_html.append(row_html)
                    table_html.append("</table></div>")
                    body_parts.append("\n".join(table_html))
                    i += consumed
                    continue

            if line.startswith("# "):
                h_text = html.escape(line[2:].strip())
                body_parts.append(f"<h1>{h_text}</h1>")
            elif line.startswith("## "):
                h_text = html.escape(line[3:].strip())
                body_parts.append(f"<h2>{h_text}</h2>")
            elif line.startswith("### "):
                h_text = html.escape(line[4:].strip())
                body_parts.append(f"<h3>{h_text}</h3>")
            else:
                is_bullet = line.startswith("- ") or line.startswith("* ")
                raw_text = re.sub(r"^[-*]\s+", "", line) if is_bullet else line
                
                # Format text
                formatted_text = html.escape(raw_text)
                formatted_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", formatted_text)
                formatted_text = re.sub(
                    r"\[(.*?)\]\((https?://[^\s\)]+)\)",
                    r'<a href="\2" target="_blank" rel="noopener noreferrer">\1</a>',
                    formatted_text,
                )
                
                # In-text citation badges
                formatted_text = re.sub(
                    r"\[(\d+)\]",
                    r'<span class="citation">[\1]</span>',
                    formatted_text,
                )

                if is_bullet:
                    body_parts.append(f"<ul><li>{formatted_text}</li></ul>")
                else:
                    body_parts.append(f"<p>{formatted_text}</p>")
            i += 1

        body_content = "\n".join(body_parts)

        html_doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>ResAI Report: {html.escape(clean_doc_title)}</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            line-height: 1.6;
            color: #1e293b;
            background-color: #f8fafc;
            margin: 0;
            padding: 40px 20px;
        }}
        .container {{
            max-width: 900px;
            margin: 0 auto;
            background: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 48px;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05);
        }}
        .badge {{
            display: inline-block;
            background-color: #dcebff;
            color: #312e81;
            font-size: 11px;
            font-weight: 700;
            padding: 4px 10px;
            border-radius: 9999px;
            text-transform: uppercase;
            letter-spacing: 0.05em;
            margin-bottom: 12px;
            border: 1px solid #bfdbfe;
        }}
        h1 {{
            font-size: 26px;
            font-weight: 800;
            color: #0f172a;
            margin-top: 0;
            margin-bottom: 24px;
            line-height: 1.3;
        }}
        h2 {{
            font-size: 18px;
            font-weight: 700;
            color: #312e81;
            margin-top: 32px;
            margin-bottom: 12px;
            padding-bottom: 6px;
            border-bottom: 1px solid #e2e8f0;
        }}
        h3 {{
            font-size: 14px;
            font-weight: 600;
            color: #475569;
            margin-top: 20px;
            margin-bottom: 8px;
        }}
        p {{
            font-size: 14px;
            color: #334155;
            margin-bottom: 14px;
        }}
        strong {{
            color: #1e293b;
            font-weight: 600;
        }}
        ul {{
            margin-top: 6px;
            margin-bottom: 14px;
            padding-left: 20px;
        }}
        li {{
            font-size: 14px;
            color: #334155;
            margin-bottom: 6px;
        }}
        a {{
            color: #635bff;
            text-decoration: none;
        }}
        a:hover {{
            text-decoration: underline;
        }}
        .citation {{
            display: inline-block;
            font-family: monospace;
            font-size: 11px;
            font-weight: bold;
            color: #312e81;
            background-color: #dcebff;
            padding: 1px 5px;
            border-radius: 4px;
            border: 1px solid #bfdbfe;
            margin: 0 2px;
        }}
        .table-container {{
            overflow-x: auto;
            margin: 20px 0;
            border: 1px solid #c4b5fd;
            border-radius: 8px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 13px;
            text-align: left;
        }}
        th {{
            background-color: #dcd6fa;
            color: #312e81;
            font-weight: 700;
            padding: 10px 14px;
            border-bottom: 1px solid #c4b5fd;
        }}
        td {{
            padding: 10px 14px;
            border-bottom: 1px solid #e2e8f0;
            color: #334155;
        }}
        tr:nth-child(even) td {{
            background-color: #f8fafc;
        }}
        tr:hover td {{
            background-color: #f1f5f9;
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="badge">ResAI Verified Evidence Report</div>
        {body_content}
    </div>
</body>
</html>"""
        return html_doc.encode("utf-8")

    @staticmethod
    def export_pdf(title: str, markdown_content: str) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            spaceAfter=12,
            textColor=colors.HexColor("#0f172a")
        )
        heading_style = ParagraphStyle(
            "ReportHeading",
            parent=styles["Heading2"],
            fontSize=12,
            leading=16,
            spaceBefore=14,
            spaceAfter=6,
            textColor=colors.HexColor("#312e81")
        )
        subheading_style = ParagraphStyle(
            "ReportSubheading",
            parent=styles["Heading3"],
            fontSize=10.5,
            leading=14,
            spaceBefore=8,
            spaceAfter=4,
            textColor=colors.HexColor("#475569")
        )
        body_style = ParagraphStyle(
            "ReportBody",
            parent=styles["Normal"],
            fontSize=9,
            leading=13,
            spaceAfter=5,
            textColor=colors.HexColor("#1e293b")
        )
        bullet_style = ParagraphStyle(
            "ReportBullet",
            parent=styles["Normal"],
            fontSize=9,
            leading=13,
            leftIndent=12,
            spaceAfter=3,
            textColor=colors.HexColor("#1e293b")
        )
        table_header_style = ParagraphStyle(
            "TableHeader",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#312e81")
        )
        table_cell_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontSize=7.5,
            leading=10,
            textColor=colors.HexColor("#1e293b")
        )

        elements = []
        clean_doc_title = title.replace("Research Report: ", "").strip()
        elements.append(Paragraph(f"<b>ResAI Research Report: {clean_doc_title}</b>", title_style))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0"), spaceBefore=2, spaceAfter=8))

        lines = markdown_content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue

            # Check for table
            if line.startswith("|") or ("|" in line and i + 1 < len(lines) and re.match(r"^\|?\s*[-:]+[-| :]*$", lines[i+1].strip())):
                table_rows, consumed = parse_markdown_table(lines[i:])
                if table_rows:
                    table_data = []
                    for r_idx, row in enumerate(table_rows):
                        row_paras = []
                        st = table_header_style if r_idx == 0 else table_cell_style
                        for cell in row:
                            safe_cell = cell.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                            safe_cell = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", safe_cell)
                            safe_cell = re.sub(r"\[(.*?)\]\((https?://[^\s\)]+)\)", r'<a href="\2" color="#635BFF"><u>\1</u></a>', safe_cell)
                            row_paras.append(Paragraph(safe_cell, st))
                        table_data.append(row_paras)

                    num_cols = len(table_rows[0])
                    available_width = 540
                    col_width = available_width / max(1, num_cols)
                    t = Table(table_data, colWidths=[col_width] * num_cols)
                    t.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcd6fa")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c4b5fd")),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                    ]))
                    elements.append(Spacer(1, 4))
                    elements.append(t)
                    elements.append(Spacer(1, 6))
                    i += consumed
                    continue

            if line.startswith("# "):
                i += 1
                continue
            elif line.startswith("## "):
                clean_h = line.lstrip("#").strip()
                elements.append(Paragraph(f"<b>{clean_h}</b>", heading_style))
            elif line.startswith("### "):
                clean_h = line.lstrip("#").strip()
                elements.append(Paragraph(f"<b>{clean_h}</b>", subheading_style))
            else:
                is_bullet = line.startswith("- ") or line.startswith("* ") or bool(re.match(r"^\d+\.\s+", line))
                raw_text = re.sub(r"^[-*]\s+|\d+\.\s+", "", line) if is_bullet else line
                
                safe_text = raw_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                formatted_text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", safe_text)
                formatted_text = re.sub(r"\[(.*?)\]\((https?://[^\s\)]+)\)", r'<a href="\2" color="#635BFF"><u>\1</u></a>', formatted_text)

                if is_bullet:
                    elements.append(Paragraph(f"&bull; {formatted_text}", bullet_style))
                else:
                    elements.append(Paragraph(formatted_text, body_style))
            i += 1

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_docx(title: str, markdown_content: str) -> bytes:
        doc = Document()
        clean_doc_title = title.replace("Research Report: ", "").strip()
        doc.add_heading(f"ResAI Research Report: {clean_doc_title}", level=1)
        
        lines = markdown_content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue

            # Check for table
            if line.startswith("|") or ("|" in line and i + 1 < len(lines) and re.match(r"^\|?\s*[-:]+[-| :]*$", lines[i+1].strip())):
                table_rows, consumed = parse_markdown_table(lines[i:])
                if table_rows:
                    table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                    table.style = 'Table Grid'
                    for r_idx, row in enumerate(table_rows):
                        for c_idx, cell in enumerate(row):
                            clean_cell = re.sub(r"\*\*(.*?)\*\*", r"\1", cell)
                            clean_cell = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", clean_cell)
                            table.cell(r_idx, c_idx).text = clean_cell
                    i += consumed
                    continue

            if line.startswith("# "):
                i += 1
                continue
            elif line.startswith("## "):
                clean_h = line.lstrip("#").strip()
                doc.add_heading(clean_h, level=2)
            elif line.startswith("### "):
                clean_h = line.lstrip("#").strip()
                doc.add_heading(clean_h, level=3)
            else:
                is_bullet = line.startswith("- ") or line.startswith("* ") or bool(re.match(r"^\d+\.\s+", line))
                raw_text = re.sub(r"^[-*]\s+|\d+\.\s+", "", line) if is_bullet else line
                clean_text = re.sub(r"\*\*(.*?)\*\*", r"\1", raw_text)
                clean_text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", clean_text)
                
                if is_bullet:
                    doc.add_paragraph(clean_text, style="List Bullet")
                else:
                    doc.add_paragraph(clean_text)
            i += 1

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


export_service = ExportService()

